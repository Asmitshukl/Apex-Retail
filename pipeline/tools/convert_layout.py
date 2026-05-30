"""Print raw workbook contents for store layout inspection."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK_PATH = Path(__file__).resolve().parents[1] / "data" / "store_layout.xlsx"


def main() -> None:
    """Print workbook rows plus sheet images, ranges, named ranges, and merges."""
    workbook = load_workbook(WORKBOOK_PATH, data_only=True)

    for sheet_name in workbook.sheetnames:
        print(sheet_name)
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(values_only=True):
            print(list(row))

        print(
            "dimensions",
            {
                "min_row": sheet.min_row,
                "max_row": sheet.max_row,
                "min_col": sheet.min_column,
                "max_col": sheet.max_column,
            },
        )

        images = getattr(sheet, "_images", [])
        print(f"images count: {len(images)}")
        for index, image in enumerate(images, start=1):
            anchor = getattr(image, "anchor", None)
            description = {
                "index": index,
                "type": type(image).__name__,
                "path": getattr(image, "path", None),
                "width": getattr(image, "width", None),
                "height": getattr(image, "height", None),
                "anchor": str(anchor),
            }
            print("image", description)

        merged_ranges = list(sheet.merged_cells.ranges)
        print(f"merged cells count: {len(merged_ranges)}")
        for merged_range in merged_ranges:
            print("merged cell", str(merged_range))

    print("named ranges")
    defined_names = list(workbook.defined_names.values())
    print(f"named ranges count: {len(defined_names)}")
    for defined_name in defined_names:
        print("named range", defined_name.name, list(defined_name.destinations))


if __name__ == "__main__":
    main()
